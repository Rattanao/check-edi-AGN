"""เว็บตรวจเอกสารขาเข้าเรือ (MANIFEST × ENTER → EDI.xlsx) สำหรับใช้ในเครื่อง/วงแลน

อัปโหลด MANIFEST (.xls/.xlsx) + ENTER (.pdf) ใส่ SHED NO. แล้วกดตรวจ — เรียก scripts/build_edi.py ตัวเดิม
แสดงผลเป็นตารางสีเดียวกับ EDI.xlsx และให้ดาวน์โหลดไฟล์ได้

รัน:  python webapp/app.py            → เปิด http://localhost:5000
      เครื่องอื่นในวงแลนเปิด http://<IP เครื่องนี้>:5000
"""
import contextlib, io, os, re, shutil, sys, threading, time, uuid

from flask import Flask, abort, render_template_string, request, send_file
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
import build_edi  # noqa: E402

JOBS = os.path.join(ROOT, 'webapp', 'jobs')
KEEP_SECONDS = 7 * 24 * 3600          # เก็บงานเก่าไว้ 7 วัน
os.makedirs(JOBS, exist_ok=True)

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024
_lock = threading.Lock()               # build_edi ใช้ตัวแปร global (EXPECTED_SHED) — ตรวจทีละงาน


def _cleanup():
    now = time.time()
    for d in os.listdir(JOBS):
        p = os.path.join(JOBS, d)
        if os.path.isdir(p) and now - os.path.getmtime(p) > KEEP_SECONDS:
            shutil.rmtree(p, ignore_errors=True)


def _read_report(path):
    """อ่าน EDI.xlsx กลับมาเป็นแถว (ข้อความ, สีพื้น) เพื่อแสดงบนเว็บให้เหมือนไฟล์"""
    ws = openpyxl.load_workbook(path).active
    title = ws.cell(2, 1).value or ''
    head = [c.value or '' for c in ws[4]]
    rows = []
    for r in ws.iter_rows(min_row=5):
        cells = []
        for c in r[:len(head)]:
            rgb = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == 'solid' else None
            bg = ('#' + rgb[-6:]) if isinstance(rgb, str) and rgb[-6:] not in ('FFFFFF', '000000') else ''
            cells.append({'v': '' if c.value is None else str(c.value), 'bg': bg})
        if any(x['v'] for x in cells):
            rows.append(cells)
    return title, head, rows


PAGE = r"""<!doctype html>
<html lang="th"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>ตรวจเอกสารขาเข้าเรือ</title>
<style>
 body{font-family:"Segoe UI",Tahoma,sans-serif;margin:0;background:#f4f6f9;color:#222}
 header{background:#1F4E78;color:#fff;padding:14px 24px;font-size:20px;font-weight:600}
 main{padding:20px 24px;max-width:1600px;margin:auto}
 .card{background:#fff;border-radius:8px;box-shadow:0 1px 3px #0002;padding:18px 20px;margin-bottom:18px}
 form{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:14px;align-items:end}
 label{font-weight:600;font-size:14px;display:block;margin-bottom:6px}
 input[type=text]{width:100%;padding:8px;font-size:15px;box-sizing:border-box;border:1px solid #bbb;border-radius:5px}
 button,.btn{background:#1F4E78;color:#fff;border:0;border-radius:5px;padding:10px 18px;font-size:15px;cursor:pointer;text-decoration:none;display:inline-block}
 .btn.green{background:#2e7d32}
 .err{background:#FDD7D7;border-left:5px solid #c00;padding:12px}
 .stat{display:flex;flex-wrap:wrap;gap:10px;margin:10px 0}
 .stat div{background:#eef3f8;border-radius:6px;padding:8px 12px;font-size:14px}
 .stat div.bad{background:#FDD7D7;font-weight:600}
 ul.crit li{margin:4px 0;color:#9c0006}
 .tbl{overflow:auto;max-height:70vh;border:1px solid #ccc}
 table{border-collapse:collapse;font-size:12px;font-family:"Aptos Narrow","Arial Narrow",Arial,sans-serif}
 th{background:#1F4E78;color:#fff;position:sticky;top:0;padding:6px;border:1px solid #fff;white-space:nowrap}
 td{border:1px solid #ddd;padding:4px 6px;vertical-align:top;max-width:280px;white-space:pre-wrap}
 td:first-child{min-width:200px}
 .muted{color:#666;font-size:13px}
</style></head><body>
<header>ตรวจเอกสารขาเข้าเรือ — MANIFEST × ENTER → EDI.xlsx</header>
<main>
<div class="card">
 <form method="post" action="/check" enctype="multipart/form-data">
  <div><label>MANIFEST (.xls / .xlsx)</label><input type="file" name="manifest" accept=".xls,.xlsx" required></div>
  <div><label>ENTER (.pdf)</label><input type="file" name="enter" accept=".pdf" required></div>
  <div><label>SHED NO. (เช่น 0141)</label><input type="text" name="shed" value="{{ shed or '' }}" placeholder="0141" required pattern="\d{3,4}"></div>
  <div><button type="submit">ตรวจสอบ</button></div>
 </form>
 <p class="muted">SHED NO.: ทุก B/L ต้องตรงกับเลขที่ใส่ (ตรง = -) ยกเว้น LAOS ใช้กฎเดิม 0124</p>
</div>
{% if error %}<div class="card err">{{ error }}</div>{% endif %}
{% if result %}
<div class="card">
 <div style="display:flex;justify-content:space-between;flex-wrap:wrap;gap:10px;align-items:center">
  <div><b>{{ result.title }}</b><div class="muted">MANIFEST: {{ result.mname }} | ENTER: {{ result.ename }} | SHED {{ shed }}</div></div>
  <a class="btn green" href="/download/{{ result.job }}">ดาวน์โหลด EDI.xlsx</a>
 </div>
 <div class="stat">
  {% for k, v, bad in result.stats %}<div class="{{ 'bad' if bad }}">{{ k }}: {{ v }}</div>{% endfor %}
 </div>
 {% if result.crit %}<b>จุดที่ต้องตรวจสอบ</b><ul class="crit">{% for c in result.crit %}<li>{{ c }}</li>{% endfor %}</ul>
 {% else %}<p style="color:#2e7d32;font-weight:600">✔ ไม่พบจุดที่ต้องแก้</p>{% endif %}
</div>
<div class="card tbl"><table>
 <tr>{% for h in result.head %}<th>{{ h }}</th>{% endfor %}</tr>
 {% for r in result.rows %}<tr>{% for c in r %}<td{% if c.bg %} style="background:{{ c.bg }}"{% endif %}>{{ c.v }}</td>{% endfor %}</tr>{% endfor %}
</table></div>
{% endif %}
</main></body></html>"""


@app.get('/')
def index():
    return render_template_string(PAGE)


@app.post('/check')
def check():
    _cleanup()
    shed = (request.form.get('shed') or '').strip()
    fm, fe = request.files.get('manifest'), request.files.get('enter')
    if not fm or not fm.filename or not fe or not fe.filename:
        return render_template_string(PAGE, error='กรุณาเลือกไฟล์ MANIFEST และ ENTER ให้ครบ', shed=shed)
    if not re.fullmatch(r'\d{3,4}', shed):
        return render_template_string(PAGE, error='SHED NO. ต้องเป็นตัวเลข 3–4 หลัก เช่น 0141', shed=shed)

    mext = os.path.splitext(fm.filename)[1].lower()
    if mext not in ('.xls', '.xlsx') or not fe.filename.lower().endswith('.pdf'):
        return render_template_string(PAGE, error='ชนิดไฟล์ไม่ถูกต้อง (MANIFEST ต้องเป็น .xls/.xlsx, ENTER ต้องเป็น .pdf)', shed=shed)
    job = uuid.uuid4().hex
    d = os.path.join(JOBS, job)
    os.makedirs(d)
    mp, ep = os.path.join(d, 'MANIFEST' + mext), os.path.join(d, 'ENTER.pdf')
    fm.save(mp)
    fe.save(ep)

    buf = io.StringIO()
    with _lock:
        try:
            with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
                build_edi.main(['--manifest', mp, '--enter', ep, '--outdir', d, '--shed', shed])
        except SystemExit as e:
            msg = str(e.code) if e.code not in (None, 0) else buf.getvalue()
            return render_template_string(PAGE, error=f'ตรวจไม่สำเร็จ: {msg}', shed=shed)
        except Exception as e:
            return render_template_string(PAGE, error=f'ตรวจไม่สำเร็จ: {type(e).__name__}: {e}', shed=shed)

    out = buf.getvalue()
    stats = []
    for line in out.splitlines():
        if line.startswith(('B/L ทั้งหมด', 'SHED NO.')):
            for part in line.split('|'):
                part = part.strip()
                mnum = re.search(r'(\d+)', part)
                label = part[:mnum.start()].strip() if mnum else part
                val = part[mnum.start():].strip() if mnum else ''
                bad = bool(mnum and int(mnum.group(1)) > 0 and not part.startswith('B/L ทั้งหมด'))
                stats.append((label, val, bad))
        elif line.startswith('⚠'):
            stats.append((line.strip(), '', True))
    crit = [l.strip()[2:] for l in out.splitlines() if l.startswith('  - ')]

    title, head, rows = _read_report(os.path.join(d, 'EDI.xlsx'))
    result = {'job': job, 'title': title, 'head': head, 'rows': rows, 'stats': stats, 'crit': crit,
              'mname': fm.filename, 'ename': fe.filename}
    return render_template_string(PAGE, result=result, shed=shed)


@app.get('/download/<job>')
def download(job):
    if not re.fullmatch(r'[0-9a-f]{32}', job):
        abort(404)
    p = os.path.join(JOBS, job, 'EDI.xlsx')
    if not os.path.exists(p):
        abort(404)
    return send_file(p, as_attachment=True, download_name='EDI.xlsx')


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print(f'เปิดเว็บที่ http://localhost:{port}  (เครื่องอื่นในวงแลน: http://<IP เครื่องนี้>:{port})')
    app.run(host='0.0.0.0', port=port, debug=False)
